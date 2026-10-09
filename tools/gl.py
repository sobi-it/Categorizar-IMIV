import openpyxl, subprocess, json, sys, os
from common import *
def run(name,changes,res_changes=None):
    wb=openpyxl.load_workbook(SRC); ws=wb['INPUTS PROYECTO']
    for k,v in changes.items(): ws[k].value=v
    for k,v in (res_changes or {}).items(): wb['RESULTADOS IMIV'][k].value=v
    p=f'/home/claude/gen/c_{name}.xlsx'; wb.save(p)
    r=subprocess.run(['python3','/mnt/skills/public/xlsx/scripts/recalc.py',p,'180'],capture_output=True,text=True)
    assert 'total_formulas' in r.stdout,r.stdout[:300]
    return openpyxl.load_workbook(p,data_only=True)
def rng(ws,r): return [[c.value for c in row] for row in ws[r]]
def results(wv):
    r=wv['RESULTADOS IMIV']
    return dict(ent=rng(r,'C21:F25'),sal=rng(r,'C28:F32'),tot=rng(r,'C35:F39'),cat=rng(r,'C4:G8'),ai=rng(r,'E12:F16'))
EDU=['Establec. Educacionales Tipo1','Establec. Educacionales Tipo2','Establec. Educacionales Tip3','Establec. Educacionales Tip 4']
def sheets_flow(wv):
    out={}
    for s in EDU: out[s]=dict(ent=rng(wv[s],'L10:O14'),sal=rng(wv[s],'L17:O21'))
    for s in ('Centros Comerciales','Supermercados-LocalesComercials','Otro Comercio (Vega)'):
        out[s]=dict(ent=rng(wv[s],'P10:S14'),sal=rng(wv[s],'P17:S21'),label_ent=wv[s]['M10'].value,label_sal=wv[s]['M17'].value)
    out['Puertos']=dict(ent=rng(wv['Puertos'],'L10:O14'),sal=rng(wv['Puertos'],'L17:O21'))
    t=wv['Terminales TPUB sin Pasajeros']
    out['Terminales TPUB sin Pasajeros']=dict(ent=rng(t,'N6:Q10'),sal=rng(t,'N13:Q17'),E6_10=[t[f'E{r}'].value for r in range(6,11)])
    return out
