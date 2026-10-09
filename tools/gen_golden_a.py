import openpyxl, subprocess, json, sys
from common import *
wb0=openpyxl.load_workbook(SRC)
T=template_sheets(wb0); assert len(T)==36
IC={s:input_cell(wb0,s) for s in T}
def run(name,changes):
    wb=openpyxl.load_workbook(SRC); ws=wb['INPUTS PROYECTO']
    for k,v in changes.items(): ws[k].value=v
    p=f'/home/claude/gen/{name}.xlsx'; wb.save(p)
    r=subprocess.run(['python3','/mnt/skills/public/xlsx/scripts/recalc.py',p,'180'],capture_output=True,text=True)
    assert 'total_formulas' in r.stdout,r.stdout[:300]
    return openpyxl.load_workbook(p,data_only=True)
def rng(ws,r): return [[c.value for c in row] for row in ws[r]]
def results(wv):
    r=wv['RESULTADOS IMIV']
    return dict(ent=rng(r,'C21:F25'),sal=rng(r,'C28:F32'),tot=rng(r,'C35:F39'),cat=rng(r,'C4:G8'),ai=rng(r,'E12:F16'))
out={}
# A) todos los módulos implementados activos
ch={}
for i,s in enumerate(T): ch[IC[s]]=round(250.5+37.25*i,2)
ch['B6']=137; ch['B10']=137; ch['B9']=137*75
wv=run('ALL',ch)
sheet={}
for s in T:
    ws=wv[s]; sheet[s]=dict(ent=rng(ws,'N8:Q12'),sal=rng(ws,'N15:Q19'))
for s in VIV:
    ws=wv[s]; sheet[s]=dict(ent=rng(ws,'M10:P14'),sal=rng(ws,'M17:P21'),c4=ws['C4'].value)
out['ALL']=dict(inputs=ch,sheets=sheet,results=results(wv))
json.dump(out,open('golden_all.json','w'))
print('ALL ok')
