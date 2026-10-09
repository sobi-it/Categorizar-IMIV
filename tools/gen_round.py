import openpyxl, subprocess, json, random
random.seed(7)
vals=[0,0.5,1.5,2.5,3.5,-0.5,-1.5,-2.5,0.49999999999999994,65.5,53.5,22.499999999999996,22.5,0.1+0.2,1.005*100,2.675*100,4.35*100,1e-12,-1e-12,19.5,20.5,79.5,80.5,249.5,250.5,39.5,40.5,159.5,160.5,499.5,500.5,999.5,1000.5]
for _ in range(180): vals.append(round(random.uniform(-50,3000),random.choice([1,2,3,4])))
for _ in range(60): vals.append(round(random.randint(0,4000)/8,3))          # múltiplos de .125 (exactos)
for _ in range(60): vals.append(random.randint(0,300)*0.07+random.randint(0,300)*0.35)  # sumas con error binario
wb=openpyxl.Workbook(); ws=wb.active
for i,v in enumerate(vals,1):
    ws.cell(i,1,v); ws.cell(i,2,f'=ROUND(A{i},0)')
wb.save('round.xlsx')
r=subprocess.run(['python3','/mnt/skills/public/xlsx/scripts/recalc.py','round.xlsx','120'],capture_output=True,text=True); print(r.stdout[:80])
w=openpyxl.load_workbook('round.xlsx',data_only=True).active
res=[[vals[i-1],w.cell(i,2).value] for i in range(1,len(vals)+1)]
json.dump(res,open('golden_round.json','w'))
print(len(res),res[:6])
