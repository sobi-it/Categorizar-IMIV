import openpyxl, re, json, os, collections
from common import *
import defs
wb=openpyxl.load_workbook(SRC); wv=openpyxl.load_workbook(SRC,data_only=True)
res=wb['RESULTADOS IMIV']
OUT='/home/claude/p001-imiv-web/js/data'; os.makedirs(OUT,exist_ok=True)
PER=['PM-L','PMD-L','PT-L','PMD-F','PT-F']
slug=lambda s:re.sub(r'[^a-z0-9]+','_',s.lower().translate(str.maketrans('áéíóúñ','aeioun'))).strip('_')
def js(path,body):
    open(f'{OUT}/{path}','w',encoding='utf8').write("/* Generado desde P001_FLUJOS_IMIV_IT_V4.0.xlsx por script (no editar a mano). */\n(function(r){'use strict';r.IMIV=r.IMIV||{};r.IMIV.model=r.IMIV.model||{};\n"+body+"\n})(typeof window!=='undefined'?window:globalThis);\n")
def dumps(o): return json.dumps(o,ensure_ascii=False,indent=1)
T=template_sheets(wb); EDU=PEND[:4]; COM=PEND[4:7]
# ================= RATES =================
rates={'plantilla':{},'viviendas':{}}
for s in T:
    ws=wb[s]; d={'sheet':s,'cells':{'ent':'E8:H12','sal':'E15:H19'},'ent':{},'sal':{}}
    for k,p in enumerate(PER):
        d['ent'][p]=[ws.cell(8+k,5+j).value for j in range(4)]; d['sal'][p]=[ws.cell(15+k,5+j).value for j in range(4)]
    assert all(isinstance(x,(int,float)) for v in list(d['ent'].values())+list(d['sal'].values()) for x in v)
    rates['plantilla'][slug(s)]=d
rx=re.compile(r"^=([\d.]+)\+([\d.]+)\*\(\$C\$4-(\d+)\)$")
for s in VIV:
    ws=wb[s]; wsv=wv[s]; tr=[]
    for t,(lo,hi) in enumerate([(None,50),(50,60),(60,140),(140,280),(280,None)]):
        e0=10+16*t; s0=17+16*t
        T_={'id':f'T{t+1}','lo':lo,'hi':hi,'cells':{'ent':f'D{e0}:G{e0+4}','sal':f'D{s0}:G{s0+4}','dummy':f'H{e0}:H{s0+4}'},'dummyEntRow':e0,'dummySalRow':s0}
        if t in (0,4):
            T_['kind']='const'; T_['ent']={p:[wsv.cell(e0+k,4+j).value for j in range(4)] for k,p in enumerate(PER)}; T_['sal']={p:[wsv.cell(s0+k,4+j).value for j in range(4)] for k,p in enumerate(PER)}
        else:
            fe=ws.cell(e0+2,3).value; fs=ws.cell(s0,3).value; me=rx.match(fe); ms=rx.match(fs); assert me and ms and me.groups()==ms.groups()
            T_['kind']='formula'; T_['base']={'a':float(me.group(1)),'b':float(me.group(2)),'x0':float(me.group(3)),'cellEnt':ws.cell(e0+2,3).coordinate,'cellSal':ws.cell(s0,3).coordinate}
            for k in range(5):
                if k!=2: assert re.match(r"^=0\.2\*\$C\$\d+$",ws.cell(e0+k,3).value)
                if k!=0: assert re.match(r"^=0\.2\*\$C\$\d+$",ws.cell(s0+k,3).value)
            T_['otherFactor']=0.2; sh=None
            for br in list(range(e0,e0+5))+list(range(s0,s0+5)):
                cur=[float(re.match(r"^=\+\$C\d+\*([\d.]+)$",ws.cell(br,4+j).value).group(1)) for j in range(4)]
                sh=sh or cur; assert cur==sh
            T_['share']=sh
        tr.append(T_)
    rates['viviendas'][slug(s)]={'sheet':s,'tramos':tr,'dummyRule':'tramo activo = (C4 > lo) y (C4 <= hi); T1: C4 <= 50; T5: C4 > 280','flowCells':{'ent':'M10:P14','sal':'M17:P21'}}
# --- educación
def tab(ws,r0,c0): return {p:[ws.cell(r0+k,c0+j).value for j in range(4)] for k,p in enumerate(PER)}
base=None
for s in EDU:
    ws=wb[s]; snap={c.coordinate:c.value for r in ws.iter_rows(min_row=7) for c in r if c.value is not None}
    if base is None: base=snap
    else: assert snap==base,s
ws=wb[EDU[0]]; titles=[ws['B7'].value,ws['B23'].value,ws['B39'].value,ws['B56'].value]
ent_s=[10,26,42,59]; sal_s=[17,33,49,66]; tipos=[]
for n in range(4):
    for r0 in (ent_s[n],sal_s[n]):
        for k in range(5): assert ws.cell(r0+k,7).value==f'=+IF($C$4={n+1},1,0)'
    tipos.append({'tipo':n+1,'label':titles[n].strip(),'ent':tab(ws,ent_s[n],3),'sal':tab(ws,sal_s[n],3),'entRow':ent_s[n],'salRow':sal_s[n]})
d94=res['D94'].value; assert re.findall(r"'([^']+)'!L10",d94)==EDU
rates['education']={'sheets':EDU,'codeCell':'C4','studentsCell':'C5','unitCell':{'ent':'J14','sal':'J21'},'dummyCol':'G','rateCols':'C:F','flowCols':'L:O','flowRows':{'ent':'L10:O14','sal':'L17:O21'},'tipos':tipos,'sumOrder':'Tipo1 + Tipo2 + Tip3 + Tip 4 (RESULTADOS IMIV!D94…)'}
# --- comercios
cr=re.compile(r"\$C\$5\s*(>=|<=|>|<)\s*(\d+)")
def rule(f):
    c=[(a,float(b)) for a,b in cr.findall(f)]; assert c; return [[a,int(b)] for a,b in c]
COMSTART={'Centros Comerciales':([10,26],[17,33]),'Supermercados-LocalesComercials':([10,26,42],[17,33,49]),'Otro Comercio (Vega)':([10,26,42,58],[17,33,49,65])}
CID={'Centros Comerciales':'centros_comerciales','Supermercados-LocalesComercials':'supermercados','Otro Comercio (Vega)':'vega'}
com={}
for s in COM:
    ws=wb[s]; es,ss=COMSTART[s]; scales=[]
    for i,(e0,s0) in enumerate(zip(es,ss)):
        rules={rule(ws.cell(r,10).value).__repr__() for r in list(range(e0,e0+5))+list(range(s0,s0+5))}; assert len(rules)==1
        scales.append({'id':f'E{i+1}','label':ws.cell(e0,3).value,'rule':rule(ws.cell(e0,10).value),'ruleText':ws.cell(e0,10).value,'ent':tab(ws,e0,6),'sal':tab(ws,s0,6),'entRow':e0,'salRow':s0})
    lab=[]
    for r in range(26,30):
        f=ws.cell(r,14).value; lab.append({'rule':None,'text':ws.cell(r,15).value,'cell':f'N{r}','formula':f})
    exp=['=+IF($C$5>L26,1,0)','=+IF(AND($C$5>L27,$C$5<M27),1,0)','=+IF(AND($C$5>L28,$C$5<M28),1,0)','=+IF(AND($C$5>L29,$C$5<M29),1,0)']
    assert [l['formula'] for l in lab]==exp and [ws[f'L{r}'].value for r in (26,27,28,29)]==[6000,1000,250,0] and ws['M27'].value=='=+L26' and ws['M28'].value=='=+L27' and ws['M29'].value==250
    # reglas de etiqueta: N26 C5>L26(6000); N27..N29 AND(C5>L,C5<M)
    labrules=[[['>',6000]],[['>',1000],['<',6000]],[['>',250],['<',1000]],[['>',0],['<',250]]]
    for l,lr in zip(lab,labrules): l['rule']=lr
    com[CID[s]]={'sheet':s,'scales':scales,'labelRules':lab,'unitCell':{'ent':'N14','sal':'N21'},'unitCellSalActual':re.findall(r"\$N\$(\d+)",ws['P17'].value)[0],'flowCols':'P:S','rateCols':'F:I','dummyCol':'J','labelCell':'M10 / M17'}
rates['commerce']=com
# --- puertos
ws=wb['Puertos']; pt=[]
for n,(e0,s0) in enumerate([(10,17),(26,33)],1):
    for r0 in (e0,s0):
        for k in range(5): assert ws.cell(r0+k,7).value==f'=+IF($C$4={n},1,0)'
    pt.append({'tipo':n,'label':ws.cell(7 if n==1 else 23,2).value.strip(),'ent':tab(ws,e0,3),'sal':tab(ws,s0,3),'entRow':e0,'salRow':s0})
rates['ports']={'sheet':'Puertos','tipos':pt,'unitCell':{'ent':'J14','sal':'J21'},'flowRows':{'ent':'L10:O14','sal':'L17:O21'}}
# --- terminal sin pasajeros
ws=wb['Terminales TPUB sin Pasajeros']
rates['terminal']={'sheet':'Terminales TPUB sin Pasajeros','factors':{'A':ws['C21'].value,'B2':ws['C22'].value,'C2':ws['C23'].value},'factorCells':{'A':'C21','B2':'C22','C2':'C23'},
 'flotaKeys':[ws.cell(r,2).value for r in range(27,33)],'freqKeys':[ws.cell(r,2).value for r in range(37,43)],
 'note':'VLOOKUP($B37,$B$27:$F$32,5,FALSE): sólo la primera coincidencia de cada clave (US1 → fila 27, US2 → fila 28).','flows':{'ent':'N6:Q10','sal':'N13:Q17'}}
open(f'{OUT}/rates.js','w',encoding='utf8').write("/* Generado desde P001_FLUJOS_IMIV_IT_V4.0.xlsx por script (no editar a mano). */\n(function(r){'use strict';r.IMIV=r.IMIV||{};r.IMIV.rates="+dumps(rates)+";})(typeof window!=='undefined'?window:globalThis);\n")
# ================= SCHEMA =================
NICE,FIELD=defs.NICE,defs.FIELD
inputs=[]
def add(id,cell,sheet,sheetCell,section,label,unit,group,module,kind='qty',**k):
    d={'id':id,'cell':cell,'sheet':sheet,'sheetCell':sheetCell,'section':section,'label':label,'unit':unit,'group':group,'module':module,'kind':kind,'enabled':True}; d.update(k); inputs.append(d)
for s in T:
    cell=input_cell(wb,s); col=re.match(r'[A-Z]+',cell).group(0); lab,unit=FIELD.get(s,defs.DEFAULT_FIELD); tid=slug(s)
    add(tid+'.cantidad',cell,s,f'{s}!C4',NICE[s],lab,unit,defs.group_of(col),'M2A',block=tid)
add('casas.viviendas','B6','Casas','Casas!C5','Casas','Cantidad de viviendas','viv.','residencial','M2B',block='casas',help='B5 = 61 × B6 (parámetro interno, no editable).')
add('departamentos.superficie_util','B9','Departamentos','Departamentos!C4 (= B9 / B10)','Departamentos','Superficie útil total','m²','residencial','M2B',block='departamentos')
add('departamentos.viviendas','B10','Departamentos','Departamentos!C5','Departamentos','Cantidad de viviendas','viv.','residencial','M2B',block='departamentos')
for i,(c1,c2) in enumerate([('H20','H21'),('H22','H23'),('H24','H25'),('H26','H27')],1):
    add(f'educacionales.tipo{i}.codigo',c1,EDU[i-1],f'{EDU[i-1]}!C4','Establecimientos educacionales',f'Establecimiento {i}: tipo','código','equipamiento','M2C','code',options=[0,1,2,3,4],block='educacionales',help='1 Parvularia/Pre-básica · 2 Básica, Media, Especial y Técnico-profesional · 3 Superior · 4 Otros · 0 = sin tipo (el Excel no aplica tasas con un código fuera de 1–4).')
    add(f'educacionales.tipo{i}.estudiantes',c2,EDU[i-1],f'{EDU[i-1]}!C5','Establecimientos educacionales',f'Establecimiento {i}: cantidad de estudiantes','estudiantes','equipamiento','M2C',block='educacionales')
for s,(sc,cc),nm in [('Centros Comerciales',('E8','E9'),'Centros comerciales - tienda por departamentos'),('Supermercados-LocalesComercials',('E12','E13'),'Supermercados - locales comerciales'),('Otro Comercio (Vega)',('E31','E32'),'Otro comercio (Vega)')]:
    b=CID[s]
    add(f'{b}.superficie',sc,s,f'{s}!C4',nm,'Superficie','m²','equipamiento','M2D',block=b)
    add(f'{b}.carga',cc,s,f'{s}!C5',nm,'Carga de ocupación','pax','equipamiento','M2D',block=b,help='Define la escala (reglas propias de cada hoja; no armonizadas con la norma).')
add('puertos.tipo','Q12','Puertos','Puertos!C4','Puertos','Tipo de puerto','código','infraestructura','M2E','code',options=[0,1,2],block='puertos',help='1 Puerto con contenedores · 2 Puerto de carga a granel · 0 = sin tipo (flujo 0 en el Excel).')
add('puertos.capacidad','Q13','Puertos','Puertos!C5','Puertos','Capacidad de transferencia frente de atraque','TEU/h o t/h','infraestructura','M2E',block='puertos')
UNIT_ROWS=['US1','US2','US1 (fila duplicada)','US2 (fila duplicada)','US1 (fila duplicada)','US2 (fila duplicada)']
for r in range(6):
    for c,(col,t) in enumerate(zip('UVW',['A','B2','C2'])):
        en=r<2
        add(f'terminal_sin_pasajeros.flota.{r+1}.{t}',f'{col}{5+r}','Terminales TPUB sin Pasajeros',f'Terminales TPUB sin Pasajeros!{"CDE"[c]}{27+r}','Terminal de buses sin pasajeros',f'Flota {UNIT_ROWS[r]} · tipo {t}','buses','infraestructura','M2F',block='terminal_sin_pasajeros',enabled=en,
            **({} if en else {'disabledReason':'Fila de flota con clave duplicada (US1/US2): la fórmula VLOOKUP del Excel nunca la utiliza.'}))
for r in range(6):
    for c,(col,p) in enumerate(zip('UVWXY',PER)):
        en=r<2
        add(f'terminal_sin_pasajeros.frec.us{r+1}.{p}',f'{col}{14+r}','Terminales TPUB sin Pasajeros',f'Terminales TPUB sin Pasajeros!{"CDEFG"[c]}{37+r}','Terminal de buses sin pasajeros',f'Frecuencia US{r+1} · {p}','buses/h','infraestructura','M2F',block='terminal_sin_pasajeros',enabled=en,
            **({} if en else {'disabledReason':'US3–US6 no están definidas en la tabla de flota del Excel (VLOOKUP devuelve #N/A). No se incluyen en el cálculo.'}))
GROUPS={'residencial':'Residencial','equipamiento':'Equipamiento','actividades':'Actividades productivas','infraestructura':'Infraestructura y terminales'}
open(f'{OUT}/inputsSchema.js','w',encoding='utf8').write("/* Generado desde P001_FLUJOS_IMIV_IT_V4.0.xlsx por script (no editar a mano). */\n(function(r){'use strict';r.IMIV=r.IMIV||{};r.IMIV.model=r.IMIV.model||{};r.IMIV.model.inputs="+dumps(inputs)+";r.IMIV.model.groups="+dumps(GROUPS)+";})(typeof window!=='undefined'?window:globalThis);\n")
# ================= MAPPINGS =================
pat=re.compile(r"(?:'([^']+)'|([A-Za-z0-9_\.]+))!")
ent_first=[46,62,78,94,110,126,142,158,174,190,206,222,238,255,272,289,306,323,340,357,374,391,408,425,442,459,476,493,510,527,544,561,578,595,612,629,646,663,680,697,714,731,748,765]
assert res['C21'].value.replace('=+','').split('+')==['D%d'%e for e in ent_first]
FAM={'Casas':'viviendas','Departamentos':'viviendas','Puertos':'ports','Terminales TPUB sin Pasajeros':'terminal','Centros Comerciales':'commerce','Supermercados-LocalesComercials':'commerce','Otro Comercio (Vega)':'commerce'}
blocks=[]
for i,e in enumerate(ent_first,1):
    sh=sorted({a or b for a,b in pat.findall(res[f'D{e}'].value)})
    if len(sh)>1: sheet,fam,bid='Establec. Educacionales','education','educacionales'
    else: sheet=sh[0]; fam=FAM.get(sheet,'template'); bid=CID.get(sheet) or slug(sheet)
    if sheet=='Terminales TPUB sin Pasajeros': bid='terminal_sin_pasajeros'
    blocks.append({'order':i,'id':bid,'sheet':sheet,'family':fam,'entRow':e,'salRow':e+7,'resultCols':'D:G'})
templates={}
for s in T:
    L12=wb[s]['L12'].value; templates[slug(s)]={'sheet':s,'inputId':slug(s)+'.cantidad','divisor':'C4/100' if '/100' in L12 else 'C4','factorPriv':2.5 if s in('Bodegas Industriales','Relleno Sanitario') else 1,'name':NICE[s],'unit':FIELD.get(s,defs.DEFAULT_FIELD)[1],'fieldLabel':FIELD.get(s,defs.DEFAULT_FIELD)[0],'unitCells':{'ent':'L12','sal':'L19'},'flowCells':{'ent':'N8:Q12','sal':'N15:Q19'}}
cellToId={f['cell']:f['id'] for f in inputs}; assert len(cellToId)==len(inputs)
# excepciones (post-procesadas: módulos ya implementados)
EXC=[]
for e in defs.EXC:
    d=dict(zip(['id','module','sheet','cell','excel','web','reason','status'],e))
    if d['id']=='EX-34': continue
    if d['module'] in ('M2C','M2D','M2E','M2F') and d['id'] not in ('EX-24','EX-25','EX-26'):
        d['status']='implementado'; d['web']=d['web'].replace('Se reproducirá','Se reproduce').replace(' (etapa M2D)','').replace('Módulo no implementado en esta etapa.','Reproducido.'); d['reason']=d['reason'].replace(' Módulo M2D no implementado en esta etapa.','').replace('Módulo no implementado en esta etapa.','Decisión aprobada.')
    EXC.append(d)
for d in EXC:
    if d['id']=='EX-24':
        d['web']='US1 y US2 (primera ocurrencia de cada clave) se calculan con la fórmula actual (frecuencia × factor medio de flota, filas 27 y 28). US3–US6 y las filas de flota duplicadas (U7:W10) quedan deshabilitadas, visibles y con advertencia; no aportan al cálculo.'
        d['reason']='Decisión D2: reproducir sólo lo verificable; no inventar el significado de las filas ni corregir referencias.'; d['status']='bloqueado'
    if d['id']=='EX-25': d['status']='implementado'; d['web']='Descartada (no alimenta cálculos).'
    if d['id']=='EX-26': d['web']='No se interpreta. Las filas se muestran tal cual, deshabilitadas, con la advertencia de D2.'
EXC+= [
 {'id':'EX-35','module':'M2F','sheet':'Terminales TPUB sin Pasajeros','cell':'E6:E17, N6:Q17','excel':'Sólo el flujo privado se alimenta (E6:E10 = H43:L43, total de veq/h); público, peatones y ciclos son constantes 0; la salida es copia de la entrada; L10 y L17 (=1) no se usan.','web':'Igual: privado = Σ(frecuencia US × factor medio); salida = entrada; otros modos = 0.','reason':'Replicar.','status':'implementado'},
 {'id':'EX-36','module':'M2C','sheet':'Establec. Educacionales (4 hojas)','cell':'L10:O21; RESULTADOS IMIV!D94','excel':'Cada hoja contiene las 4 tablas de tipo; un dummy G elige la tabla según C4; RESULTADOS suma las 4 hojas (D94 = Tipo1 + Tipo2 + Tip3 + Tip 4).','web':'Cuatro filas de establecimiento (código + estudiantes). La selección usa una lista cerrada 0–4; el motor acepta cualquier código (modo auditoría/pruebas).','reason':'La lista cerrada evita códigos inválidos en uso normal; el motor reproduce el Excel para cualquier valor.','status':'implementado'},
 {'id':'EX-37','module':'M2D','sheet':'Comercios','cell':'M10, M17, N26:O29','excel':'La etiqueta de escala usa límites estrictos (>6000, (1000;6000), (250;1000), (0;250)) y es independiente de las dummies J; con carga 6.000, 1.000, 250 ó 0 devuelve #N/A aunque las dummies sí activen una escala.','web':'Se calcula la etiqueta con las mismas reglas; si no hay escala se muestra «Sin escala aplicable según la lógica actual del Excel.». Las dummies y los flujos no se alteran.','reason':'Decisión D4.','status':'implementado'},
]
DIF=[dict(zip(['id','topic','excel','web','type','ref'],d)) for d in defs.DIF]
for d in DIF:
    if d['id']=='DIF-05': d['web']='«Sin escala aplicable según la lógica actual del Excel.»'
    if d['id']=='DIF-06': d['web']='US3–US6 y filas de flota duplicadas deshabilitadas con advertencia; US1–US2 se calculan como el Excel.'
DIF=[d for d in DIF if d['id']!='DIF-10']
DIF+= [{'id':'DIF-13','topic':'Modo de operación','excel':'Un único modo.','web':'Modo usuario (bloquea texto y negativos) y modo auditoría/pruebas (permite negativos para comparar). El texto se bloquea en ambos modos.','type':'Intencional (D6)','ref':'EX-13, EX-14'},
 {'id':'DIF-14','topic':'Códigos de tipo (educacionales, puertos)','excel':'Celda libre sin validación.','web':'Lista cerrada en la interfaz; el motor reproduce cualquier código.','type':'Entrada de datos','ref':'EX-36'},
 {'id':'DIF-15','topic':'Terminal de buses sin pasajeros: US3–US6 en pruebas','excel':'#N/A en todos los resultados si hay frecuencia en US3–US6.','web':'Se excluyen del cálculo y se informa. El caso se registra como diferencia intencional (D2).','type':'Intencional (D2)','ref':'EX-24'}]
model_map={'version':'1.0.0','source':'P001_FLUJOS_IMIV_IT_V4.0.xlsx',
 'periods':[{'id':'PM-L','excel':'PM-L'},{'id':'PMD-L','excel':'PMD-L (INPUTS: PMd-L)'},{'id':'PT-L','excel':'PT-L'},{'id':'PMD-F','excel':'PMD-F (INPUTS: PMd-F)'},{'id':'PT-F','excel':'PT-F'}],
 'modes':[{'id':'priv','label':'Transporte privado','short':'Privado','unit':'veh/h'},{'id':'pub','label':'Transporte público','short':'T. público','unit':'viajes/h'},{'id':'peat','label':'Peatones','short':'Peatones','unit':'viajes/h'},{'id':'cic','label':'Ciclos','short':'Ciclos','unit':'viajes/h'}],
 'constants':{'CASAS_SUP_UTIL_M2_POR_VIVIENDA':{'value':61,'source':'INPUTS PROYECTO!B5 = 61*B6'}},
 'thresholds':{'catPriv':{'lo':20,'mid':80,'hi':250,'cells':'RESULTADOS IMIV!E4:E8'},'catOtros':{'lo':40,'mid':160,'hi':500,'cells':'RESULTADOS IMIV!F4:F8'},
  'aiPriv':{'edges':[80,120,180,250,400,550,750,1000],'values':[2,3,4,5,6,7,8],'over':'Art 3.2.2.','cells':'RESULTADOS IMIV!E12:E16'},'aiOtros':{'edges':[160,500,2000],'values':[3,6,8],'cells':'RESULTADOS IMIV!F12:F16'}},
 'blocks':blocks,'templates':templates,'cellToId':cellToId,
 'resultCells':{'ent':'C21:F25','sal':'C28:F32','tot':'C35:F39','cat':'C4:G8','ai':'C12:F16'},
 'notes':{'casas':'Para esta versión, la superficie útil de casas se calcula como 61 m² por vivienda, reproduciendo la fórmula original del Excel.','art322':'El libro original no implementa el desglose de intersecciones posteriores. El resultado se conserva como referencia al Art. 3.2.2.','us36':'US3–US6 no se encuentran correctamente definidos en el libro Excel original y quedan pendientes de confirmación. No se incluyen en el cálculo de esta versión.','escalas':'Los límites de escala de esta versión reproducen el Excel y aún no han sido armonizados con la normativa.','sinEscala':'Sin escala aplicable según la lógica actual del Excel.'},
 'exceptions':EXC,'differences':DIF}
open(f'{OUT}/mappings.js','w',encoding='utf8').write("/* Generado desde P001_FLUJOS_IMIV_IT_V4.0.xlsx por script (no editar a mano). */\n(function(r){'use strict';r.IMIV=r.IMIV||{};r.IMIV.model=Object.assign(r.IMIV.model||{},"+dumps(model_map)+");})(typeof window!=='undefined'?window:globalThis);\n")
json.dump({'exceptions':EXC,'differences':DIF,'blocks':blocks,'inputs':inputs,'templates':templates,'rates':rates},open('project_data.json','w'),ensure_ascii=False)
print(len(inputs),len(blocks),collections.Counter(b['family'] for b in blocks),len(EXC),len(DIF))
