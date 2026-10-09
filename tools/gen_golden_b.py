import openpyxl, subprocess, json, sys
from common import *
def run(name,changes):
    wb=openpyxl.load_workbook(SRC); ws=wb['INPUTS PROYECTO']
    for k,v in changes.items(): ws[k].value=v
    p=f'/home/claude/gen/{name}.xlsx'; wb.save(p)
    r=subprocess.run(['python3','/mnt/skills/public/xlsx/scripts/recalc.py',p,'180'],capture_output=True,text=True)
    assert 'total_formulas' in r.stdout,r.stdout[:300]
    return openpyxl.load_workbook(p,data_only=True)
def rng(ws,r): return [[c.value for c in row] for row in ws[r]]
AVG=[0,45,50,50.5,55,60,60.5,100,140,140.5,200,280,280.5,300]
out={}
for a in AVG:
    viv=200
    wv=run(f'S_{a}',{'B5':a*viv,'B6':viv,'B9':a*viv,'B10':viv})
    d={}
    for s in ('Casas','Departamentos'):
        ws=wv[s]; d[s]=dict(c4=ws['C4'].value,ent=rng(ws,'M10:P14'),sal=rng(ws,'M17:P21'))
    out[str(a)]=dict(viviendas=viv,sup_total=a*viv,sheets=d)
    print(a,flush=True)
json.dump(out,open('golden_tramos.json','w'))
